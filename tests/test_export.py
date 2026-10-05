"""The Excel export must open in Excel (checked here with openpyxl)."""

import io

import openpyxl

from conftest import embedding


def test_class_workbook_opens(fresh_db):
    from app.data.export import class_workbook

    db = fresh_db
    cls = db.create_class("IT 301 / Systems: Analysis?", "BSIT 3A")  # sheet-name-hostile
    ana = db.add_student("2023-0001", "Ana Cruz", embedding(1), cls["id"])
    sess = db.create_session("Week 1", cls["id"])
    db.set_student_status(sess, ana, "Late")
    db.end_session(sess)

    filename, data = class_workbook(cls["id"])
    assert filename.endswith(".xlsx")
    wb = openpyxl.load_workbook(io.BytesIO(data))
    assert len(wb.sheetnames) >= 2
    text = [str(c.value) for ws in wb for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Ana Cruz" in text and "2023-0001" in text


def test_missing_class_has_no_workbook(fresh_db):
    from app.data.export import class_workbook

    assert class_workbook(999) is None


def test_safe_sheet_name():
    from app.data.xlsx import safe_sheet_name

    name = safe_sheet_name("A/B\\C?D*E[F]G:" + "x" * 40)
    assert len(name) <= 31
    assert not set(name) & set("/\\?*[]:")
